from pathlib import Path
import datetime as _datetime
import subprocess
import uuid
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn


OUT_DIR = Path("/Users/zds/bsr/tmp/moegambit_patent")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_DOCX = Path("/Users/zds/Desktop/一种面向稀疏混合专家模型训练的运行时混合恢复方法及系统_专利申请.docx")

TITLE = "一种面向稀疏混合专家模型训练的运行时混合恢复方法及系统"


def set_east_asia_font(run, east="宋体", latin="Times New Roman", size=12, bold=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_format(p, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_line:
        pf.first_line_indent = Pt(24)


def add_para(doc, text="", first_line=True, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=first_line, align=align)
    r = p.add_run(text)
    set_east_asia_font(r, size=size, bold=bold)
    return p


def add_center(doc, text, size=14, bold=True):
    return add_para(doc, text, first_line=False, bold=bold, align=WD_ALIGN_PARAGRAPH.CENTER, size=size)


def add_heading_plain(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    set_east_asia_font(r, size=15, bold=True)
    return p


def add_page_break(doc):
    doc.add_page_break()


abstract = (
    "本发明公开了一种面向稀疏混合专家模型训练的运行时混合恢复方法及系统，属于大模型分布式训练与容错恢复技术领域。"
    "该方法在训练运行过程中接收失效rank事件，基于安全点控制器阻止故障期间的不完整优化器提交；"
    "将训练状态划分为密集数据并行复制的非专家状态、专家并行分片的专家状态以及运行时元数据；"
    "根据故障rank、故障步、最近检查点步和专家集合计算检查点间隔、窗口内专家陈旧债务以及专家加权陈旧密度；"
    "在健康密集数据并行peer可用且陈旧密度满足阈值约束时，从健康peer拉取当前非专家状态，并从检查点或专家peer恢复rank本地专家状态；"
    "进一步通过权重优先、优化器稍后的两阶段协议在更新屏障保护下恢复训练。"
    "本发明能够在不回滚整个训练作业的情况下对混合专家模型的rank失效进行状态感知恢复，减少检查点读写和重放迭代开销，并通过结构化日志保证恢复决策可审计。"
)


claims = [
    "一种面向稀疏混合专家模型训练的运行时混合恢复方法，其特征在于，包括：在包含流水线并行、密集张量并行、专家并行和专家张量并行的分布式混合专家模型训练过程中，接收失效rank事件，所述失效rank事件表示为\\(\\langle r,t,c\\rangle\\)，其中\\(r\\)为失效逻辑rank，\\(t\\)为安全点处理后应继续执行的训练步，\\(c\\)为最近一次可用检查点对应的训练步；在所述失效rank事件发生后，将失效rank以及失效rank对应的专家集合标记为恢复中状态，并在优化器提交点安装提交保护，以阻止使用恢复期间产生的中间梯度执行参数更新；将所述失效rank的训练状态划分为密集数据并行复制的非专家状态、专家并行分片的专家状态以及由rank布局派生的运行时元数据；根据\\(\\Delta=t-c\\)计算检查点间隔，并基于窗口内历史混合恢复事件维护专家陈旧债务\\(S(t)\\)；根据所述专家陈旧债务、失效rank对应的新增专家集合、专家总数以及暴露窗口长度计算专家加权陈旧密度\\(\\Phi'(t)\\)；根据健康密集数据并行peer是否可用、检查点间隔是否位于预设范围以及专家加权陈旧密度是否不超过预设阈值，确定采用混合恢复路径或检查点重启路径；当采用混合恢复路径时，从健康密集数据并行peer同步当前训练步的非专家状态，并恢复失效rank对应的专家状态；在恢复后重建运行时元数据并将替换rank重集成为健康状态。",
    "根据权利要求1所述的方法，其特征在于，所述密集数据并行复制的非专家状态包括注意力层参数、嵌入层参数、归一化层参数、路由器参数以及其对应的密集数据并行复制优化器状态；所述专家并行分片的专家状态包括专家权重以及专家权重对应的优化器一阶矩和二阶矩。",
    "根据权利要求1所述的方法，其特征在于，专家数据并行度\\(\\mathrm{EDP}\\)按照如下方式确定：\\(\\mathrm{EDP}=\\frac{W}{\\mathrm{PP}\\times\\mathrm{EP}\\times\\mathrm{ETP}}\\)，其中\\(W\\)为训练作业总rank数，\\(\\mathrm{PP}\\)为流水线并行度，\\(\\mathrm{EP}\\)为专家并行度，\\(\\mathrm{ETP}\\)为专家张量并行度；当\\(\\mathrm{EDP}=1\\)时，失效rank对应的专家状态不存在健康peer副本；当\\(\\mathrm{EDP}>1\\)时，失效rank对应的专家状态能够从健康专家peer恢复。",
    "根据权利要求1所述的方法，其特征在于，所述提交保护包括：在检测到失效事件后，将当前飞行中迭代标记为丢弃状态；任一rank到达优化器提交点时，若检测到所述丢弃状态或恢复中状态，则跳过本次优化器更新，并丢弃对应的中间梯度。",
    "根据权利要求1所述的方法，其特征在于，所述专家陈旧债务\\(S(t)\\)表示为\\(S(t)=\\sum_h |E_h|\\Delta_h\\)，其中求和范围为暴露窗口\\([t-W_{\\mathrm{exp}},t)\\)内的历史混合恢复事件，\\(E_h\\)为历史恢复事件\\(h\\)中被检查点恢复的专家集合，\\(\\Delta_h\\)为历史恢复事件\\(h\\)对应的检查点间隔，\\(W_{\\mathrm{exp}}\\)为暴露窗口长度。",
    "根据权利要求5所述的方法，其特征在于，所述专家加权陈旧密度\\(\\Phi'(t)\\)按照如下方式计算：\\(\\Phi'(t)=\\frac{S(t)+|E_{\\mathrm{new}}|\\Delta}{N_{\\mathrm{expert}}W_{\\mathrm{exp}}}\\)，其中\\(E_{\\mathrm{new}}\\)为当前失效rank对应的新增专家集合，\\(N_{\\mathrm{expert}}\\)为暴露域内专家总数，\\(\\Phi'(t)\\)表示当前事件若采用混合恢复后相对于全窗口专家暴露预算的归一化陈旧程度。",
    "根据权利要求1所述的方法，其特征在于，采用混合恢复路径的条件包括：失效rank在相同流水线并行坐标和密集张量并行坐标下存在健康密集数据并行peer；\\(\\Delta\\)不小于预设最小间隔阈值\\(\\Delta_{\\min}\\)；\\(\\Delta\\)不大于预设最大单事件间隔阈值\\(\\Delta_{\\max}\\)；并且\\(\\Phi'(t)\\)不大于预设陈旧密度阈值\\(\\Phi_{\\max}\\)；若上述任一条件不满足，则采用检查点重启路径。",
    "根据权利要求1所述的方法，其特征在于，所述混合恢复路径包括并行执行的第一状态路径和第二状态路径；所述第一状态路径从健康密集数据并行peer拉取当前训练步的非专家状态；所述第二状态路径在\\(\\mathrm{EDP}=1\\)时从最近检查点的rank本地分片读取专家权重和专家优化器状态，在\\(\\mathrm{EDP}>1\\)且专家peer可用时从健康专家peer拉取专家状态。",
    "根据权利要求1所述的方法，其特征在于，恢复专家状态时采用权重优先、优化器稍后的两阶段协议：第一阶段恢复非专家状态和专家权重，使替换rank能够进入前向传播和反向传播；第二阶段在后台恢复专家优化器状态，并在专家优化器状态未附着前通过更新屏障阻止受影响专家的优化器提交。",
    "根据权利要求1所述的方法，其特征在于，所述重建运行时元数据包括重建专家目录、刷新调度拓扑、重建通信进程组以及更新失效rank到替换rank的映射关系，并将替换rank依次置于恢复中、已修复、屏障等待和健康状态。",
    "根据权利要求1所述的方法，其特征在于，在每次恢复事件中生成结构化恢复日志，所述结构化恢复日志至少包括检查点间隔\\(\\Delta\\)、新增专家集合大小\\(|E_{\\mathrm{new}}|\\)、窗口专家陈旧债务\\(S(t)\\)、专家加权陈旧密度\\(\\Phi'(t)\\)、恢复路径决策、fallback原因、各恢复阶段时延以及重集成状态迁移记录。",
    "一种实现权利要求1-11任一项所述方法的面向稀疏混合专家模型训练的运行时混合恢复系统，其特征在于，包括：故障检测与替换rank管理模块，用于接收或生成失效rank事件并分配替换rank；安全点控制模块，用于标记恢复中状态并安装优化器提交保护；状态分类模块，用于区分非专家状态、专家状态和运行时元数据；恢复策略判定模块，用于计算\\(\\Delta\\)、\\(S(t)\\)和\\(\\Phi'(t)\\)，并基于peer可用性和阈值约束选择混合恢复路径或检查点重启路径；状态恢复模块，用于从健康peer和检查点分别恢复对应状态；两阶段恢复模块，用于执行权重优先和优化器稍后的恢复过程；重集成与日志模块，用于重建元数据、释放更新屏障并生成结构化恢复日志。",
    "根据权利要求12所述的系统，其特征在于，所述恢复策略判定模块维护一滑动暴露窗口，并以\\(O(1)\\)方式在历史混合恢复事件进入或离开所述滑动暴露窗口时更新专家陈旧债务\\(S(t)\\)。",
    "根据权利要求12所述的系统，其特征在于，所述状态恢复模块包括peer拉取子模块和分片读取子模块；所述peer拉取子模块用于在相同流水线并行坐标和密集张量并行坐标下从健康dense-DP peer同步当前非专家状态；所述分片读取子模块用于从检查点读取失效rank对应的专家分片。",
    "一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1-11任一项所述的方法。"
]


description_sections = [
    ("技术领域", [
        "本发明属于大模型分布式训练、混合专家模型训练和运行时容错恢复技术领域，具体涉及一种面向稀疏混合专家模型训练的运行时混合恢复方法及系统，尤其适用于多机多卡GPU集群中采用流水线并行、密集张量并行、专家并行和专家张量并行进行大规模预训练或继续训练的场景。"
    ]),
    ("背景技术", [
        "随着大规模语言模型和生成式人工智能技术的发展，模型参数规模、训练数据规模和训练持续时间持续增加。为了在有限显存和通信带宽条件下训练超大规模模型，现有训练系统通常综合采用数据并行、张量并行、流水线并行以及专家并行等混合并行策略。其中，稀疏混合专家模型通过在每个token上仅激活部分专家，在扩大总参数量的同时控制单次计算量，已经成为大模型训练的重要形态。",
        "在大规模训练运行过程中，rank失效、GPU故障、节点重启、通信超时等事件难以完全避免。现有主流恢复方式通常采用检查点重启，即从最近一次全局一致检查点恢复训练状态，并重放从检查点步到故障步之间的全部迭代。该方式虽然语义简单，但会丢弃检查点之后已经完成的GPU计算，并产生较大的状态加载和迭代重放开销。",
        "对于密集模型训练而言，如果失效rank的全部模型状态在数据并行维度上存在健康副本，则可以通过从健康peer拉取状态的方式减少检查点重启开销。然而，稀疏混合专家模型的训练状态具有异构性：非专家层和路由器等状态通常在密集数据并行维度上复制，而专家权重及其优化器状态则沿专家并行维度或专家张量并行维度进行分片。在专家数据并行度\\(\\mathrm{EDP}=1\\)的布局下，失效rank拥有的专家分片在当前内存中不存在健康副本。",
        "因此，现有检查点重启方法未能利用混合专家训练状态的异构恢复来源，往往将可从健康peer恢复的非专家状态和必须从检查点恢复的专家状态作为一个整体回滚；而现有密集peer恢复方法又假设失效rank的全部状态均存在健康副本，难以直接用于\\(\\mathrm{EDP}=1\\)的混合专家训练场景。",
        "此外，若仅将非专家状态恢复到当前步而将专家状态恢复到检查点步，则会引入专家状态相对于非专家状态的陈旧性。该陈旧性并非必然导致训练失败，但如果在多个rank或多个时间窗口内累积，可能影响专家负载均衡和训练轨迹。现有技术缺乏一种能够在运行时量化该陈旧性、据此判断是否允许混合恢复、并对恢复过程进行结构化审计的技术方案。"
    ]),
    ("发明内容", [
        "针对上述问题，本发明提出一种面向稀疏混合专家模型训练的运行时混合恢复方法及系统，旨在解决现有检查点重启恢复开销大、密集peer恢复无法覆盖\\(\\mathrm{EDP}=1\\)专家分片、以及部分恢复缺乏质量风险约束和审计机制的问题。",
        "本发明的核心思想在于：将混合专家训练状态按照恢复来源划分为非专家复制态、专家分片态和运行时元数据；在rank失效时优先从健康密集数据并行peer恢复当前非专家状态，仅对缺乏健康副本的专家状态执行检查点分片恢复；同时通过专家加权陈旧密度\\(\\Phi'(t)\\)约束专家状态陈旧暴露，并通过安全点控制、两阶段恢复和结构化日志保证恢复语义正确和可审计。",
        "为实现上述目的，本发明提供一种运行时混合恢复方法，包括以下步骤：",
        "步骤一，训练状态建模。对于总rank数为\\(W\\)的混合专家训练作业，按照流水线并行度\\(\\mathrm{PP}\\)、专家并行度\\(\\mathrm{EP}\\)、专家张量并行度\\(\\mathrm{ETP}\\)确定专家数据并行度\\(\\mathrm{EDP}=\\frac{W}{\\mathrm{PP}\\times\\mathrm{EP}\\times\\mathrm{ETP}}\\)。在每个流水线阶段和密集张量并行坐标内，将注意力层、嵌入层、归一化层和路由器参数划分为密集数据并行复制的非专家状态，将专家权重及对应优化器状态划分为专家并行分片的专家状态，将专家目录、调度拓扑和进程组信息划分为运行时元数据。",
        "步骤二，故障事件接收与安全点建立。当训练作业检测到失效rank事件\\(\\langle r,t,c\\rangle\\)时，将失效rank及其拥有的专家集合标记为恢复中，并在优化器提交点安装提交保护；若故障发生在训练迭代中间，则丢弃当前飞行中迭代，使后续恢复从合法训练步\\(t\\)继续执行。",
        "步骤三，陈旧暴露计算。计算检查点间隔\\(\\Delta=t-c\\)，维护暴露窗口\\([t-W_{\\mathrm{exp}},t)\\)内历史混合恢复事件的专家陈旧债务\\(S(t)=\\sum_h |E_h|\\Delta_h\\)，并根据\\(\\Phi'(t)=\\frac{S(t)+|E_{\\mathrm{new}}|\\Delta}{N_{\\mathrm{expert}}W_{\\mathrm{exp}}}\\)计算当前事件若采用混合恢复时的专家加权陈旧密度。",
        "步骤四，恢复路径判定。当健康密集数据并行peer可用、\\(\\Delta\\)位于预设阈值范围\\([\\Delta_{\\min},\\Delta_{\\max}]\\)内且\\(\\Phi'(t)\\le\\Phi_{\\max}\\)时，选择混合恢复路径；否则选择检查点重启路径。所述阈值可根据检查点间隔、恢复成本和训练质量噪声水平进行预设或校准。",
        "步骤五，混合状态恢复。对于非专家状态，从相同流水线并行坐标和密集张量并行坐标下的健康dense-DP peer拉取当前训练步状态；对于专家状态，在\\(\\mathrm{EDP}=1\\)时从最近检查点的rank本地专家分片恢复，在\\(\\mathrm{EDP}>1\\)且专家peer可用时从健康专家peer恢复；对于运行时元数据，重新计算专家目录、调度拓扑和进程组映射。",
        "步骤六，两阶段恢复。第一阶段先恢复非专家状态和专家权重，使替换rank能够尽快参与前向传播和反向传播；第二阶段在后台恢复专家优化器状态，并在优化器状态未附着前通过更新屏障阻止受影响专家提交不完整更新。",
        "步骤七，重集成与审计。替换rank依次经过恢复中、已修复、屏障等待和健康状态，恢复控制器生成结构化日志，记录恢复路径、阈值输入、fallback原因、时延分段和状态迁移，以支持离线审计和质量关联分析。",
        "本发明还提供一种实现上述方法的运行时混合恢复系统，包括故障检测与替换rank管理模块、安全点控制模块、状态分类模块、恢复策略判定模块、状态恢复模块、两阶段恢复模块以及重集成与日志模块。",
        "与现有技术相比，本发明至少具有以下技术效果：第一，避免将全部训练状态整体回滚，减少检查点加载和重放迭代开销；第二，针对\\(\\mathrm{EDP}=1\\)下专家状态无peer副本的特殊性，允许非专家状态peer恢复与专家状态检查点恢复并存；第三，通过专家加权陈旧密度对混合恢复风险进行运行时约束，避免无界累积陈旧专家状态；第四，通过安全点控制和两阶段恢复避免不完整优化器提交；第五，通过结构化日志提高恢复过程的可审计性和可复现性。"
    ]),
    ("附图说明", [
        "为了更清楚地说明本申请实施例中的技术方案，下面将对实施例描述中所需要使用的附图作简要介绍。显而易见地，下面描述中的附图仅为本申请的一些实施例，本领域普通技术人员在不付出创造性劳动的前提下，还可以根据这些附图获得其他附图。",
        "图1：运行时混合恢复系统架构图，展示故障检测、恢复控制器、策略判定、peer恢复、检查点恢复、两阶段协议和结构化日志之间的关系；",
        "图2：运行时混合恢复方法流程图，展示从故障事件接收、安全点建立、陈旧密度计算、策略判定到恢复执行和重集成的处理流程；",
        "图3：混合恢复与两阶段协议时序图，展示替换rank、健康dense-DP peer、检查点存储和全局更新屏障之间的交互关系。"
    ]),
    ("具体实施方式", [
        "下面结合附图和具体实施例对本发明作进一步说明。应理解，以下实施例仅用于说明本发明，而非用于限定本发明的保护范围。在不冲突的情况下，以下实施例中的特征可以相互组合。",
        "在一个实施例中，本发明部署于多机多卡GPU集群的大模型训练平台中，训练框架支持流水线并行、密集张量并行、专家并行、专家张量并行和密集数据并行。以包含128个路由专家的稀疏混合专家模型为例，当训练作业具有\\(W=64\\)个rank、\\(\\mathrm{PP}=8\\)、\\(\\mathrm{EP}=8\\)、\\(\\mathrm{ETP}=1\\)时，\\(\\mathrm{EDP}=\\frac{W}{\\mathrm{PP}\\times\\mathrm{EP}\\times\\mathrm{ETP}}=1\\)。此时每个流水线阶段内每个专家分片在当前内存中仅由一个rank持有，若该rank失效，则对应专家状态无法从健康专家peer直接取得。",
        "参见图1，运行时混合恢复系统嵌入训练作业内部。故障检测与替换rank管理模块可从集群监控、训练框架watchdog或进程退出信号接收失效rank事件\\(\\langle r,t,c\\rangle\\)。恢复控制器在接收事件后将失效rank标记为RECOVERING，同时通知各rank安装优化器提交保护。恢复策略判定模块读取失效rank的并行坐标、最近检查点步、专家目录和历史恢复窗口，并输出Hybrid或Restart决策。",
        "对于非专家状态，本实施例在相同\\(\\mathrm{PP}\\)和密集\\(\\mathrm{TP}\\)坐标下选择一个健康dense-DP peer。该peer持有当前训练步\\(t\\)的注意力层、嵌入层、归一化层和路由器状态，因此能够为替换rank提供最新非专家状态。对于专家状态，当\\(\\mathrm{EDP}=1\\)时，替换rank从检查点步\\(c\\)对应的rank本地专家分片中读取专家权重和优化器状态；当\\(\\mathrm{EDP}>1\\)且专家peer可用时，系统也可以从健康专家peer读取专家状态，从而进一步降低检查点读取开销。",
        "参见图2，本实施例的恢复方法首先建立安全点。若故障发生在某一训练迭代中间，则该迭代的中间梯度可能不完整或与恢复后的参数状态不一致，因此恢复控制器将该迭代标记为DISCARD。任一rank到达optimizer.step()之前均检查该标记，若处于DISCARD或RECOVERING状态，则跳过本次优化器提交。由此可避免部分rank提交了故障前梯度而部分rank使用恢复后状态的情况。",
        "恢复策略判定模块维护窗口专家陈旧债务\\(S(t)\\)。对于窗口内每一次历史混合恢复事件\\(h\\)，记被检查点恢复的专家集合为\\(E_h\\)，检查点间隔为\\(\\Delta_h\\)，则该事件对窗口债务贡献\\(|E_h|\\Delta_h\\)。窗口滑动时，进入窗口的事件贡献被加入\\(S(t)\\)，离开窗口的事件贡献从\\(S(t)\\)中扣除，从而以常数时间维护\\(S(t)\\)。当前事件的新增专家集合记为\\(E_{\\mathrm{new}}\\)，检查点间隔为\\(\\Delta=t-c\\)，则当前事件采用混合恢复后的专家加权陈旧密度为\\(\\Phi'(t)=\\frac{S(t)+|E_{\\mathrm{new}}|\\Delta}{N_{\\mathrm{expert}}W_{\\mathrm{exp}}}\\)。",
        "在一个具体策略配置中，系统设置\\(\\Delta_{\\min}=1\\)、\\(\\Delta_{\\max}=200\\)、\\(W_{\\mathrm{exp}}=2000\\)、\\(\\Phi_{\\max}=0.1\\)。当健康dense-DP peer不存在、\\(\\Delta<\\Delta_{\\min}\\)、\\(\\Delta>\\Delta_{\\max}\\)或\\(\\Phi'(t)>\\Phi_{\\max}\\)时，系统选择检查点重启路径；否则选择混合恢复路径。上述数值仅为一种实施例，实际系统可根据训练模型、检查点周期、恢复时延、网络状态和训练质量噪声水平进行调整。",
        "采用混合恢复路径时，状态恢复模块并行执行两条路径。路径P从健康dense-DP peer同步非专家状态，所述同步可以通过点对点通信、广播或其他集合通信接口实现。路径C从检查点存储读取失效rank对应的专家分片，所述专家分片包括专家权重以及对应的优化器一阶矩和二阶矩。在路径P和路径C完成专家权重恢复后，替换rank即可进入权重优先阶段。",
        "参见图3，两阶段恢复协议将权重恢复和优化器状态恢复解耦。第一阶段恢复非专家状态和专家权重，使替换rank能够尽快参与下一次前向传播和反向传播；第二阶段继续在后台读取专家优化器状态。在优化器状态尚未恢复完成前，系统通过更新屏障阻止受影响专家的优化器提交，未受影响参数可根据实现策略继续参与计算或等待屏障统一释放。当优化器状态附着完成后，屏障释放，缓冲梯度或后续梯度按照AdamW等优化器规则执行更新。",
        "重集成过程中，系统重新生成专家目录，刷新all-to-all调度拓扑，重建失效rank相关的通信进程组，并更新失效rank到替换rank的逻辑映射。替换rank依次经过RECOVERING、REPAIRED、BARRIER和HEALTHY状态。只有在状态机进入HEALTHY后，恢复控制器才允许后续训练步按正常路径执行。",
        "为便于审计，每次恢复事件都会生成结构化日志。日志内容包括失效rank标识、故障步\\(t\\)、检查点步\\(c\\)、\\(\\Delta\\)、\\(|E_{\\mathrm{new}}|\\)、\\(S(t)\\)、\\(\\Phi'(t)\\)、PeerAvail结果、阈值判定结果、选择的恢复路径、fallback原因、路径P耗时、路径C耗时、权重优先阶段耗时、优化器恢复耗时、屏障释放时间以及状态机迁移时间。该日志能够将每一次恢复决策与后续训练质量指标对应起来，便于系统调试、策略校准和合规审计。",
        "在另一个实施例中，若多个rank同时失效，系统可以按照预设顺序或并行批处理方式对各失效rank分别计算新增专家集合和检查点间隔，并将各rank对应的\\(|E_{\\mathrm{new}}|\\Delta\\)贡献纳入窗口债务。若累计\\(\\Phi'(t)\\)超过阈值，则后续失效rank或整个批次可以回退到检查点重启路径，以避免多rank突发故障导致专家陈旧暴露过大。",
        "在又一个实施例中，本发明可与现有检查点优化系统结合使用。检查点优化系统降低检查点保存和读取成本，本发明则在恢复侧判断哪些状态可以从健康peer取得、哪些状态必须从检查点取得。当混合恢复被允许时，训练作业可直接从故障步t继续，而无需重放从c到t之间的迭代；当混合恢复不被允许时，系统仍保留标准检查点重启路径，从而保证恢复策略具有保守fallback。",
        "本领域技术人员可以理解，本发明所述模块可以以软件、硬件或软硬件结合方式实现；所述方法可以由训练框架插件、运行时控制器、分布式通信库扩展或集群管理组件执行。上述实施例中使用的具体模型规模、并行度、阈值和通信实现方式均不构成对本发明保护范围的限制。凡在本发明精神和原则之内所作的等同替换、修改或组合，均应包含在本发明的保护范围之内。"
    ]),
]


def _mx_model(page_width, page_height):
    model = ET.Element(
        "mxGraphModel",
        {
            "dx": "1400",
            "dy": "900",
            "grid": "1",
            "gridSize": "10",
            "guides": "1",
            "tooltips": "1",
            "connect": "1",
            "arrows": "1",
            "fold": "1",
            "page": "1",
            "pageScale": "1",
            "pageWidth": str(page_width),
            "pageHeight": str(page_height),
            "math": "1",
            "shadow": "0",
        },
    )
    root = ET.SubElement(model, "root")
    ET.SubElement(root, "mxCell", {"id": "0"})
    ET.SubElement(root, "mxCell", {"id": "1", "parent": "0"})
    return model, root


def _vertex(root, cid, value, x, y, w, h, style):
    cell = ET.SubElement(
        root,
        "mxCell",
        {
            "id": cid,
            "value": value,
            "style": style,
            "vertex": "1",
            "parent": "1",
        },
    )
    ET.SubElement(
        cell,
        "mxGeometry",
        {
            "x": str(x),
            "y": str(y),
            "width": str(w),
            "height": str(h),
            "as": "geometry",
        },
    )
    return cell


def _edge(root, cid, source, target, label="", style=None):
    edge_style = style or (
        "edgeStyle=orthogonalEdgeStyle;rounded=1;orthogonalLoop=1;jettySize=auto;"
        "html=1;strokeColor=#333333;strokeWidth=1.4;endArrow=block;endFill=1;"
        "fontFamily=Arial;fontSize=13;"
    )
    cell = ET.SubElement(
        root,
        "mxCell",
        {
            "id": cid,
            "value": label,
            "style": edge_style,
            "edge": "1",
            "parent": "1",
            "source": source,
            "target": target,
        },
    )
    ET.SubElement(cell, "mxGeometry", {"relative": "1", "as": "geometry"})
    return cell


def _edge_abs(root, cid, x1, y1, x2, y2, label="", style=None):
    edge_style = style or (
        "endArrow=block;endFill=1;html=1;rounded=0;strokeColor=#333333;"
        "strokeWidth=1.3;fontFamily=Arial;fontSize=12;"
    )
    cell = ET.SubElement(
        root,
        "mxCell",
        {
            "id": cid,
            "value": label,
            "style": edge_style,
            "edge": "1",
            "parent": "1",
        },
    )
    geo = ET.SubElement(cell, "mxGeometry", {"relative": "1", "as": "geometry"})
    ET.SubElement(geo, "mxPoint", {"x": str(x1), "y": str(y1), "as": "sourcePoint"})
    ET.SubElement(geo, "mxPoint", {"x": str(x2), "y": str(y2), "as": "targetPoint"})
    return cell


def _text_style(size=14):
    return (
        "text;html=1;strokeColor=none;fillColor=none;align=center;verticalAlign=middle;"
        f"whiteSpace=wrap;rounded=0;fontFamily=Arial;fontSize={size};fontColor=#111111;"
    )


def _box_style(fill="#FFFFFF", stroke="#222222", size=14, bold=False):
    font_style = "1" if bold else "0"
    return (
        "rounded=1;whiteSpace=wrap;html=1;arcSize=10;absoluteArcSize=1;"
        f"fillColor={fill};strokeColor={stroke};strokeWidth=1.4;"
        f"fontFamily=Arial;fontSize={size};fontStyle={font_style};fontColor=#111111;"
        "align=center;verticalAlign=middle;spacing=10;"
    )


def _diamond_style(fill="#FFFFFF", size=14):
    return (
        "rhombus;whiteSpace=wrap;html=1;fillColor="
        f"{fill};strokeColor=#222222;strokeWidth=1.4;fontFamily=Arial;fontSize={size};"
        "fontColor=#111111;align=center;verticalAlign=middle;spacing=8;"
    )


def _write_drawio(model, path, name):
    mxfile = ET.Element(
        "mxfile",
        {
            "host": "app.diagrams.net",
            "modified": _datetime.datetime.now().isoformat(timespec="seconds"),
            "agent": "Codex",
            "version": "30.2.6",
            "type": "device",
        },
    )
    diagram = ET.SubElement(mxfile, "diagram", {"id": uuid.uuid4().hex, "name": name})
    diagram.append(model)
    tree = ET.ElementTree(mxfile)
    tree.write(path, encoding="utf-8", xml_declaration=True)


def _export_drawio(src, dst):
    drawio = "/Applications/draw.io.app/Contents/MacOS/draw.io"
    subprocess.run(
        [drawio, "-x", "-f", "png", "--scale", "2", "-b", "16", "-o", str(dst), str(src)],
        check=True,
    )


def make_figures():
    """Create editable draw.io figures and export them as PNGs for DOCX embedding."""
    figs = []

    # Fig. 1: system architecture.  A patent-style module diagram with
    # one top-level decision path and a separate hybrid-state recovery path.
    model, root = _mx_model(1280, 720)
    module = _box_style("#FFFFFF", size=13)
    main = _box_style("#F3F6F8", size=13, bold=True)
    source = _box_style("#FFFFFF", size=13)
    decision = _diamond_style("#FFFFFF", 13)
    note = _box_style("#FFFFFF", size=12)

    _vertex(root, "f1_event", "失效rank事件<br>\\(\\langle r,t,c\\rangle\\)", 40, 60, 190, 76, module)
    _vertex(root, "f1_ctrl", "恢复控制器（101）<br>安全点控制 · 提交保护", 295, 52, 250, 92, main)
    _vertex(root, "f1_policy", "恢复策略判定（102）<br>\\(\\mathrm{PeerAvail}\\land\\Delta\\in[\\Delta_{\\min},\\Delta_{\\max}]\\land\\Phi'(t)\\le\\Phi_{\\max}\\)", 610, 46, 350, 104, main)
    _vertex(root, "f1_decision", "Hybrid<br>或<br>Restart", 1015, 58, 130, 96, decision)
    _vertex(root, "f1_restart", "检查点重启路径（109）<br>全局一致恢复", 1010, 205, 200, 78, module)

    _vertex(root, "f1_class", "状态分类模块（103）<br>非专家复制态<br>专家分片态<br>运行时元数据", 40, 345, 210, 120, module)
    _vertex(root, "f1_peer", "状态来源A（104）<br>健康dense-DP peer<br>当前非专家状态", 315, 295, 240, 94, source)
    _vertex(root, "f1_expert", "状态来源B（105）<br>\\(\\mathrm{EDP}=1\\)：检查点分片<br>\\(\\mathrm{EDP}>1\\)：专家peer", 315, 445, 240, 104, source)
    _vertex(root, "f1_recover", "混合状态恢复模块（106）<br>Path P + Path C<br>重构替换rank状态", 625, 355, 260, 118, main)
    _vertex(root, "f1_phase", "两阶段恢复模块（107）<br>权重优先<br>优化器稍后", 960, 335, 230, 102, module)
    _vertex(root, "f1_log", "重集成与日志模块（108）<br>状态机：RECOVERING → HEALTHY<br>记录决策、阈值、时延", 960, 505, 240, 112, module)
    _vertex(root, "f1_label", "混合恢复路径", 640, 300, 220, 34, _text_style(14))

    arrow = "endArrow=block;endFill=1;html=1;rounded=0;strokeColor=#222222;strokeWidth=1.4;fontFamily=Arial;fontSize=12;"
    dashed_arrow = arrow + "dashed=1;"
    _edge_abs(root, "f1_e1", 230, 98, 295, 98, "", arrow)
    _edge_abs(root, "f1_e2", 545, 98, 610, 98, "", arrow)
    _edge_abs(root, "f1_e3", 960, 98, 1015, 98, "", arrow)
    _edge_abs(root, "f1_e4", 1080, 154, 1080, 205, "Restart", arrow)
    _edge_abs(root, "f1_e5", 145, 136, 145, 345, "状态建模", dashed_arrow)
    _edge_abs(root, "f1_e6", 250, 405, 315, 345, "", arrow)
    _edge_abs(root, "f1_e6b", 250, 405, 315, 497, "", arrow)
    _edge_abs(root, "f1_e7", 555, 342, 625, 392, "Path P", arrow)
    _edge_abs(root, "f1_e8", 555, 497, 625, 438, "Path C", arrow)
    _edge_abs(root, "f1_e9", 885, 414, 960, 386, "", arrow)
    _edge_abs(root, "f1_e10", 1075, 437, 1075, 505, "", arrow)
    _edge_abs(root, "f1_e11", 785, 150, 755, 355, "Hybrid", dashed_arrow)

    d1 = OUT_DIR / "fig1_system.drawio"
    p1 = OUT_DIR / "fig1_system.png"
    _write_drawio(model, d1, "图1")
    _export_drawio(d1, p1)
    figs.append(p1)

    # Fig. 2: method flow.  Left side is the common pre-check path; the
    # right side is the admitted hybrid path; the lower branch is fallback.
    model, root = _mx_model(1080, 980)
    step = _box_style("#FFFFFF", size=13)
    hybrid = _box_style("#F3F6F8", size=13)
    terminal = _box_style("#F3F6F8", size=13, bold=True)
    branch = _diamond_style("#FFFFFF", 12)

    _vertex(root, "f2_s1", "S1 接收故障事件<br>\\(\\langle r,t,c\\rangle\\)", 60, 40, 260, 76, step)
    _vertex(root, "f2_s2", "S2 建立安全点<br>丢弃飞行中迭代", 60, 155, 260, 80, step)
    _vertex(root, "f2_s3", "S3 划分训练状态<br>非专家复制态 / 专家分片态 / 元数据", 60, 275, 260, 92, step)
    _vertex(root, "f2_s4", "S4 计算恢复风险<br>\\(\\Delta=t-c\\), \\(S(t)\\), \\(\\Phi'(t)\\)", 60, 415, 260, 86, step)
    _vertex(root, "f2_decide", "是否满足<br>\\(\\mathrm{PeerAvail}\\)<br>\\(\\Delta\\)阈值<br>\\(\\Phi'(t)\\le\\Phi_{\\max}\\)", 395, 388, 178, 178, branch)

    _vertex(root, "f2_p", "S5a Path P<br>从健康dense-DP peer<br>拉取当前非专家状态", 705, 250, 270, 92, hybrid)
    _vertex(root, "f2_c", "S5b Path C<br>恢复专家状态<br>\\(\\mathrm{EDP}=1\\)：检查点分片", 705, 390, 270, 104, hybrid)
    _vertex(root, "f2_phase", "S6 两阶段恢复<br>权重优先，优化器稍后<br>更新屏障保护提交", 705, 545, 270, 104, hybrid)
    _vertex(root, "f2_rebuild", "S7 重建运行时元数据<br>通信组、专家目录、rank映射", 705, 700, 270, 92, hybrid)
    _vertex(root, "f2_done", "S8 替换rank进入HEALTHY<br>恢复训练", 705, 845, 270, 80, terminal)
    _vertex(root, "f2_restart", "S5c 检查点重启<br>任一保护条件不满足", 355, 690, 260, 86, step)

    _edge_abs(root, "f2_e1", 190, 116, 190, 155, "", arrow)
    _edge_abs(root, "f2_e2", 190, 235, 190, 275, "", arrow)
    _edge_abs(root, "f2_e3", 190, 367, 190, 415, "", arrow)
    _edge_abs(root, "f2_e4", 320, 458, 395, 477, "", arrow)
    _edge_abs(root, "f2_e5", 573, 477, 705, 296, "是", arrow)
    _edge_abs(root, "f2_e6", 840, 342, 840, 390, "", arrow)
    _edge_abs(root, "f2_e7", 840, 494, 840, 545, "", arrow)
    _edge_abs(root, "f2_e8", 840, 649, 840, 700, "", arrow)
    _edge_abs(root, "f2_e9", 840, 792, 840, 845, "", arrow)
    _edge_abs(root, "f2_e10", 484, 566, 484, 690, "否", arrow)
    _edge_abs(root, "f2_e11", 615, 733, 705, 745, "", dashed_arrow)

    d2 = OUT_DIR / "fig2_flow.drawio"
    p2 = OUT_DIR / "fig2_flow.png"
    _write_drawio(model, d2, "图2")
    _export_drawio(d2, p2)
    figs.append(p2)

    # Fig. 3: two-phase sequence protocol.
    model, root = _mx_model(1240, 760)
    actor = _box_style("#F7F8FA", size=13, bold=True)
    message = (
        "endArrow=block;endFill=1;html=1;rounded=0;strokeColor=#333333;"
        "strokeWidth=1.3;endArrow=block;endFill=1;fontFamily=Arial;fontSize=12;"
    )
    dashed = (
        "endArrow=block;endFill=1;html=1;rounded=0;strokeColor=#333333;"
        "strokeWidth=1.2;dashed=1;endArrow=block;endFill=1;fontFamily=Arial;fontSize=12;"
    )
    lifeline = "shape=rect;html=1;rounded=0;fillColor=#C8CDD2;strokeColor=none;"
    band = "rounded=0;whiteSpace=wrap;html=1;fillColor=#F6F6F6;strokeColor=#DDDDDD;strokeWidth=1;fontFamily=Arial;fontSize=12;fontColor=#333333;align=left;verticalAlign=top;spacing=8;"

    _vertex(root, "f3_band1", "阶段一：权重优先恢复", 40, 210, 1160, 160, band)
    _vertex(root, "f3_band2", "阶段二：优化器稍后恢复", 40, 370, 1160, 230, band)

    actors = [
        ("f3_ctrl", "恢复控制器", 60),
        ("f3_rank", "替换rank", 290),
        ("f3_peer", "健康dense-DP peer", 520),
        ("f3_ckpt", "检查点/专家peer", 760),
        ("f3_barrier", "更新屏障", 1010),
    ]
    for cid, label, x in actors:
        _vertex(root, cid, label, x, 40, 165, 54, actor)
        _vertex(root, cid + "_line", "", x + 82, 110, 2, 520, lifeline)

    x_ctrl, x_rank, x_peer, x_ckpt, x_barrier = 142, 372, 602, 842, 1092
    _edge_abs(root, "f3_e1", x_ctrl, 145, x_barrier, 145, "安装提交保护", message)
    _edge_abs(root, "f3_e2", x_ctrl, 185, x_rank, 185, "分配替换rank", message)
    _edge_abs(root, "f3_e3", x_rank, 250, x_peer, 250, "拉取当前非专家状态", message)
    _edge_abs(root, "f3_e4", x_rank, 300, x_ckpt, 300, "读取专家权重分片", message)
    _edge_abs(root, "f3_e5", x_ckpt, 345, x_rank, 345, "返回专家权重", dashed)
    _edge_abs(root, "f3_e6", x_rank, 415, x_barrier, 415, "进入更新屏障", message)
    _edge_abs(root, "f3_e7", x_rank, 470, x_ckpt, 470, "后台读取优化器状态", message)
    _edge_abs(root, "f3_e8", x_barrier, 520, x_rank, 520, "阻止受影响专家提交", dashed)
    _edge_abs(root, "f3_e9", x_ckpt, 565, x_rank, 565, "返回优化器状态", dashed)
    _edge_abs(root, "f3_e10", x_barrier, 615, x_rank, 615, "释放屏障并允许更新", message)
    _vertex(
        root,
        "f3_note",
        "说明：非专家状态取自当前健康peer；专家状态按检查点分片或专家peer恢复；优化器状态附着前不允许受影响专家提交。",
        120,
        665,
        980,
        54,
        _box_style("#FFFFFF", size=12),
    )

    d3 = OUT_DIR / "fig3_sequence.drawio"
    p3 = OUT_DIR / "fig3_sequence.png"
    _write_drawio(model, d3, "图3")
    _export_drawio(d3, p3)
    figs.append(p3)

    return figs


def build_doc():
    figs = make_figures()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(0.2)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5

    add_heading_plain(doc, "摘要")
    add_para(doc, abstract)
    add_page_break(doc)

    add_heading_plain(doc, "权利要求书")
    for i, claim in enumerate(claims, 1):
        add_para(doc, f"{i}. {claim}")
    add_page_break(doc)

    add_heading_plain(doc, "说明书")
    add_heading_plain(doc, TITLE)
    for heading, paras in description_sections:
        add_heading_plain(doc, heading)
        for para in paras:
            if para.startswith("步骤") or para.startswith("图") or para.startswith("本发明") or para.startswith("与现有"):
                add_para(doc, para)
            else:
                add_para(doc, para)

    add_page_break(doc)
    for idx, fig in enumerate(figs, 1):
        add_center(doc, f"图{idx}", size=14, bold=True)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(fig), width=Cm(16.0))
        if idx != len(figs):
            add_page_break(doc)

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build_doc()
    print(path)
